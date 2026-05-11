"""v5_writer — openpyxl serialization for v5 5-sheet workbook.

This module is the ONLY place that touches openpyxl. It receives:
  - subject_curr: dict from v5_core.normalize_subject(rc)
  - subject_prev: dict (optional) for comparison mode
  - a_rows / b_rows / c_rows: List[FindingRow]
and produces a styled .xlsx file at the requested path.
"""
from __future__ import annotations
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Optional

from v5_core import FindingRow, VERIFY_REQ

# --- styling ----------------------------------------------------------
THIN   = Side(border_style='thin', color='B0B0B0')
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

TITLE_FILL  = PatternFill('solid', fgColor='305496')
TITLE_FONT  = Font(bold=True, color='FFFFFF', name='맑은 고딕', size=12)
HDR_FILL    = PatternFill('solid', fgColor='4472C4')
HDR_FONT    = Font(bold=True, color='FFFFFF', name='맑은 고딕', size=10)
SECT_FILL   = PatternFill('solid', fgColor='D9E1F2')
DIVIDER_FILL= PatternFill('solid', fgColor='F2F2F2')
DIVIDER_FONT= Font(italic=True, name='맑은 고딕', size=10, color='1F4E78')

PRESENCE_FILL = {
    'present':      PatternFill('solid', fgColor='C6EFCE'),
    'absent':       PatternFill('solid', fgColor='DDEBF7'),
    'suspected':    PatternFill('solid', fgColor='FFEB9C'),
    'not_measured': PatternFill('solid', fgColor='E7E6E6'),
}
WRAP_TOP = Alignment(wrap_text=True, vertical='top')
CENTER   = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _apply_header(ws, row, headers, fill=HDR_FILL, font=HDR_FONT):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = fill; c.font = font; c.alignment = CENTER; c.border = BORDER


def _apply_divider(ws, row, text, n_cols=18):
    c = ws.cell(row=row, column=1, value=text)
    c.fill = DIVIDER_FILL; c.font = DIVIDER_FONT; c.alignment = WRAP_TOP
    for col in range(2, n_cols + 1):
        ws.cell(row=row, column=col).fill = DIVIDER_FILL


def _write_finding_rows(ws, start_row: int, rows: List[FindingRow], dividers_by_category: dict) -> int:
    """Write FindingRow list with category divider rows interleaved.
    Returns the next free row index."""
    r = start_row
    seen_cats = set()
    for fr in rows:
        cat = fr.finding_category
        if cat not in seen_cats:
            seen_cats.add(cat)
            divider_text = dividers_by_category.get(cat, f'  ▶  {cat}')
            _apply_divider(ws, r, divider_text)
            r += 1
        for ci, val in enumerate(fr.as_list(), start=1):
            cell = ws.cell(row=r, column=ci, value=val)
            cell.border = BORDER
            cell.alignment = WRAP_TOP
        # color the presence_status cell
        ps_cell = ws.cell(row=r, column=9)
        if fr.presence_status in PRESENCE_FILL:
            ps_cell.fill = PRESENCE_FILL[fr.presence_status]
        r += 1
    return r


# --- 00 sheet ---------------------------------------------------------
def _write_00(wb, curr: dict, prev: Optional[dict]):
    ws = wb.create_sheet('00_환자기본정보')
    ws['A1'] = '환자 기본정보 및 검진 측정값 요약'
    ws['A1'].fill = TITLE_FILL; ws['A1'].font = TITLE_FONT; ws['A1'].alignment = CENTER
    ws.merge_cells('A1:E1')

    headers = ['항목',
               curr.get('examinationDate', ''),
               (prev.get('examinationDate', '') if prev else ''),
               '단위', '비고']
    _apply_header(ws, 3, headers)

    sections = [
        ('■ 환자 식별 정보', [
            ('patientName (환자명)',       'patientName',       '', '이름'),
            ('birthDate (생년월일)',       'birthDate',         '', 'JSON 추정'),
            ('patient_id (비식별)',        'patient_id',        '', '연구용 ID'),
        ]),
        ('■ 검진 메타', [
            ('examinationDate (검진일)',   'examinationDate',   '', ''),
            ('screeningDate (판정일)',     'screeningDate',     '', ''),
            ('organizationName (검진기관)','organizationName',  '', ''),
            ('판정의사',                   'diagnosingDoctor',  '', ''),
        ]),
        ('■ 계측검사 — 신체 계측', [
            ('height (신장)',              'height',            'cm',     ''),
            ('weight (체중)',              'weight',            'kg',     ''),
            ('waistCircumference (허리둘레)','waistCircumference','cm',   '복부비만 기준 남≥90 여≥85'),
            ('bmi (체질량지수)',           'bmi',               'kg/m²',  '과체중 25–29.9 / 비만 ≥30'),
            ('waist 판정',                 'waistCircumference_judgment', '', ''),
            ('bmi 판정',                   'bmi_judgment',      '', ''),
        ]),
        ('■ 활력징후', [
            ('visionLeft (시력 좌)',       'visionLeft',        '', ''),
            ('visionRight (시력 우)',      'visionRight',       '', ''),
            ('hearingLeft (청력 좌)',      'hearingLeft',       '', ''),
            ('hearingRight (청력 우)',     'hearingRight',      '', ''),
            ('bloodPressure (혈압)',       'bloodPressure',     'mmHg', ''),
            ('systolic (수축기)',          'systolic',          'mmHg', ''),
            ('diastolic (이완기)',         'diastolic',         'mmHg', ''),
        ]),
        ('■ 혈액검사', [
            ('hemoglobin (혈색소)',        'hemoglobin',        'g/dL',  '여 12.0–15.5 / 남 13.0–17.5'),
            ('fastingGlucose (공복혈당)',  'fastingGlucose',    'mg/dL', 'IFG 100–125, 당뇨≥126'),
            ('totalCholesterol (총콜레스테롤)','totalCholesterol','mg/dL','정상<200'),
            ('hdlCholesterol (HDL)',       'hdlCholesterol',    'mg/dL', '여성권장≥50'),
            ('ldlCholesterol (LDL)',       'ldlCholesterol',    'mg/dL', '정상<130'),
            ('triglycerides (중성지방)',   'triglycerides',     'mg/dL', '경계 150–199'),
            ('serumCreatinine (혈청크레아티닌)','serumCreatinine','mg/dL','여성≤1.2'),
            ('egfr (신사구체여과율)',      'egfr',              'mL/min/1.73m²','정상≥60'),
            ('ast (AST/SGOT)',             'ast',               'IU/L',  '정상≤40'),
            ('alt (ALT/SGPT)',             'alt',               'IU/L',  '여성정상≤35'),
            ('gamaGtp (감마GTP)',          'gamaGtp',           'IU/L',  '여성정상≤35'),
        ]),
        ('■ 요검사 / 영상검사', [
            ('urinaryProtein (요단백)',    'urinaryProtein',    '', ''),
            ('chestXray (흉부촬영)',       'chestXray',         '', ''),
        ]),
    ]
    r = 4
    for header, items in sections:
        c = ws.cell(row=r, column=1, value=header)
        c.fill = SECT_FILL; c.font = Font(bold=True, name='맑은 고딕', size=10, color='1F4E78')
        for col in range(2, 6): ws.cell(row=r, column=col).fill = SECT_FILL
        r += 1
        for label, key, unit, note in items:
            ws.cell(row=r, column=1, value=label).font = Font(bold=True, name='맑은 고딕', size=10)
            ws.cell(row=r, column=2, value=curr.get(key))
            ws.cell(row=r, column=3, value=(prev.get(key) if prev else None))
            ws.cell(row=r, column=4, value=unit)
            ws.cell(row=r, column=5, value=note)
            for col in range(1, 6):
                ws.cell(row=r, column=col).border = BORDER
                ws.cell(row=r, column=col).alignment = WRAP_TOP
            r += 1

    ws.column_dimensions['A'].width = 32
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 32
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = 'A4'


# --- A / B / C sheets -------------------------------------------------
HEADERS_18 = [
    '#', 'patient_id', 'exam_date', 'finding_source', 'finding_category',
    'item_name_kr', 'snomed_code', 'snomed_display', 'presence_status',
    'presence_qualifier_code', 'verification_status', 'snomed_domain',
    'derived_value', 'ref_range', 'medication_related', 'mapping_confidence',
    'loinc_code', 'mapper_note',
]

A_DIVIDERS = {
    'A_계측':       '  ▶  A_계측',
    'B_혈액':       '  ▶  B_혈액',
    'C_요검사':     '  ▶  C_요검사',
    'D_영상':       '  ▶  D_영상',
    'E_생활습관':   '  ▶  E_생활습관',
    'F_문진':       '  ▶  F_문진',
    'G_특수':       '  ▶  G_특수',
    'H_질환소견':   '  ▶  H_질환소견',
    'I_암검진':     '  ▶  I_암검진 파생',
}
B_DIVIDERS = {
    'A_계측':       '  ▶  A_계측 — 임상 해석',
    'B_혈액':       '  ▶  B_혈액 — 임상 해석',
    'C_요검사':     '  ▶  C_요검사 — 임상 해석',
    'D_영상':       '  ▶  D_영상 — 임상 해석',
    'E_생활습관':   '  ▶  E_생활습관 — 임상 해석 (모두 포함)',
    'G_특수':       '  ▶  G_특수 — 항상 생성 (B형간염/골밀도/우울증/인지/노인기능/예방접종)',
    'I_암검진':     '  ▶  I_암검진 — 임상 해석 (항상 생성)',
}
C_DIVIDERS = {
    'A_계측':       '  ▶  A_계측 차이',
    'B_혈액':       '  ▶  B_혈액 차이',
    'I_암검진':     '  ▶  암검진 차이',
}

LEGEND_TEXT = ('presence_status: ■present(녹)=소견있음  ■absent(청)=없음/정상  '
               '■suspected(황)=의심  ■not_measured(회)=미시행/비해당')
B_PRINCIPLES = ('【원칙】①양성소견→해당상태 present 직접표기  ②상위진단확정→파생진단생략  '
                '③음성확인이임상적의미있는경우만absent  ④not_measured=미시행/비해당  ⑤항상생성항목표시')


def _write_finding_sheet(wb, name: str, title: str, legend: str,
                         principles: Optional[str], rows: List[FindingRow],
                         dividers: dict):
    ws = wb.create_sheet(name)
    # title
    ws['A1'] = title
    ws['A1'].fill = TITLE_FILL; ws['A1'].font = TITLE_FONT; ws['A1'].alignment = WRAP_TOP
    ws.merge_cells('A1:R1')
    # legend
    ws['A2'] = legend
    ws['A2'].font = Font(italic=True, name='맑은 고딕', size=10)
    ws.merge_cells('A2:R2')
    # principles or header
    if principles:
        ws['A3'] = principles
        ws['A3'].font = Font(name='맑은 고딕', size=10, color='1F4E78')
        ws.merge_cells('A3:R3')
        header_row = 4
    else:
        header_row = 3
    _apply_header(ws, header_row, HEADERS_18)
    next_row = _write_finding_rows(ws, header_row + 1, rows, dividers)

    # column widths
    widths = [4, 11, 12, 18, 12, 36, 12, 32, 13, 14, 13, 18, 32, 22, 8, 9, 11, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = f'A{header_row + 1}'


# --- D sheet ----------------------------------------------------------
def _write_d(wb, used_codes: set):
    ws = wb.create_sheet('D_SNOMED_코드_레퍼런스')
    ws['A1'] = 'SNOMED CT 코드 레퍼런스 — v5 신규 추가 코드 및 검증 상태'
    ws['A1'].fill = TITLE_FILL; ws['A1'].font = TITLE_FONT; ws['A1'].alignment = CENTER
    ws.merge_cells('A1:F1')
    headers = ['SNOMED 코드', '개념명 (영문)', '한국어 설명', '검증 상태', '사용 시트', '비고']
    _apply_header(ws, 3, headers)

    # Hard-coded reference set — codes that v5 introduced + ones we use
    entries = [
        ('28315006',  'Low high density lipoprotein cholesterol level', 'HDL 저하',           '✅ MCP 이전 검증',             'A+B', ''),
        ('302870006', 'Hypertriglyceridemia',                            '고중성지방혈증',     '✅ MCP 이전 검증',             'A+B', ''),
        ('166832000', 'Raised low density lipoprotein cholesterol level','LDL 상승',           '⚠ MCP 재가동 후 검증 필요',   'B',   'SNOMED 참조 코드'),
        ('406583007', 'Hepatitis B surface antigen positive',            'HBsAg 양성',         '⚠ MCP 재가동 후 검증 필요',   'A',   '비해당→not_measured'),
        ('406584001', 'Hepatitis B surface antigen negative',            'HBsAg 음성',         '⚠ MCP 재가동 후 검증 필요',   'A',   '참조용'),
        ('395699004', 'Hepatitis B surface antibody positive',           'HBsAb 양성',         '⚠ MCP 재가동 후 검증 필요',   'A',   '비해당→not_measured'),
        ('395700008', 'Hepatitis B surface antibody negative',           'HBsAb 음성',         '⚠ MCP 재가동 후 검증 필요',   'A',   '참조용'),
        ('129839007', 'At risk of falling',                              '낙상 위험',          '⚠ MCP 재가동 후 검증 필요',   'B',   '노인기능검사 비해당→not_measured'),
        ('160685001', 'Independent in activities of daily living',       'ADL 독립',           '⚠ MCP 재가동 후 검증 필요',   'A',   ''),
        ('160674001', 'Needs assistance with activities of daily living','ADL 보조 필요',      '⚠ MCP 재가동 후 검증 필요',   'B',   ''),
        ('160676004', 'Dependent for activities of daily living',        'ADL 의존',           '⚠ MCP 재가동 후 검증 필요',   'B',   ''),
        ('45850009',  'Continent of urine',                              '배뇨 정상',          '✅ 본 세션 내 MCP 검증',       'B',   ''),
        ('165232002', 'Urinary incontinence',                            '요실금',             '✅ 참조 코드',                 'A+B', ''),
        ('43994002',  'Lack of physical activity',                       '신체활동 부족',      '⚠ MCP 503 미검증 (사용자 지정)','A+B','신체활동=유산소'),
        ('40979000',  'Lack of exercise',                                '운동 부족',          '✅ MCP 이전 검증',             'A+B', '근력운동=저항운동'),
        ('86198006',  'Influenza vaccination',                           '인플루엔자 백신',    '✅ 참조 코드',                 'A+B', '항상 생성'),
        ('12866006',  'Pneumococcal vaccination',                        '폐렴구균 백신',      '✅ 참조 코드',                 'A+B', '항상 생성'),
    ]
    r = 4
    for entry in entries:
        for ci, v in enumerate(entry, start=1):
            ws.cell(row=r, column=ci, value=v).border = BORDER
            ws.cell(row=r, column=ci).alignment = WRAP_TOP
        r += 1

    widths = [16, 50, 22, 26, 14, 32]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'A4'


# --- public entry -----------------------------------------------------
def write_workbook(out_path: str,
                   curr: dict,
                   prev: Optional[dict],
                   a_rows: List[FindingRow],
                   b_rows: List[FindingRow],
                   c_rows: List[FindingRow]):
    wb = Workbook()
    # Drop default sheet
    default_ws = wb.active
    wb.remove(default_ws)

    _write_00(wb, curr, prev)

    _write_finding_sheet(wb, 'A_hc_finding',
        title='A_hc_finding — 검사·측정 직접 소견 (raw findings)',
        legend=LEGEND_TEXT, principles=None, rows=a_rows, dividers=A_DIVIDERS)

    _write_finding_sheet(wb, 'B_hc_clinical',
        title='B_hc_clinical — 임상 해석 결과 (disease presence/absence judgements)',
        legend=LEGEND_TEXT, principles=B_PRINCIPLES, rows=b_rows, dividers=B_DIVIDERS)

    title_c = (f'C_2023_비교 — {curr.get("patientName","")} {prev.get("examinationDate","") if prev else ""} '
               f'| ★ {curr.get("examinationDate","")} 대비 차이 항목')
    _write_finding_sheet(wb, 'C_2023_비교',
        title=title_c, legend=LEGEND_TEXT, principles=None, rows=c_rows, dividers=C_DIVIDERS)

    used = {fr.snomed_code for fr in a_rows + b_rows + c_rows if fr.snomed_code}
    _write_d(wb, used)

    wb.save(out_path)
    return out_path
