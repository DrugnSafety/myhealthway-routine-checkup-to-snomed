#!/usr/bin/env python3
"""verify_workbook — Schema validator for v5 .xlsx output.

Checks:
  1. Exactly 5 sheets, in order: 00_환자기본정보, A_hc_finding, B_hc_clinical,
     C_2023_비교, D_SNOMED_코드_레퍼런스
  2. A and B sheets have the 18-column header in the correct row
  3. Every data row in A/B has a non-empty snomed_code AND presence_status ∈
     {present, absent, suspected, not_measured}
  4. presence_qualifier_code matches presence_status (per RULES.md)
  5. always-generate items appear in B sheet
  6. (optional) Compare against a golden reference workbook by row count per sheet

Usage:
    python verify_workbook.py PATH.xlsx [--golden REF.xlsx] [--strict]

Exit code 0 = pass, 1 = warnings, 2 = errors.
"""
from __future__ import annotations
import argparse, json, sys, os
from openpyxl import load_workbook

EXPECTED_SHEETS = ['00_환자기본정보', 'A_hc_finding', 'B_hc_clinical',
                   'C_2023_비교', 'D_SNOMED_코드_레퍼런스']
EXPECTED_HEADERS_18 = [
    '#', 'patient_id', 'exam_date', 'finding_source', 'finding_category',
    'item_name_kr', 'snomed_code', 'snomed_display', 'presence_status',
    'presence_qualifier_code', 'verification_status', 'snomed_domain',
    'derived_value', 'ref_range', 'medication_related', 'mapping_confidence',
    'loinc_code', 'mapper_note',
]
QUAL_MAP = {'present': '410515003', 'absent': '410516002'}
ALLOWED_PS = {'present', 'absent', 'suspected', 'not_measured'}

ALWAYS_GENERATE_KEYWORDS = [
    'B형간염', '골다공증', '우울증', '인지기능', '낙상', 'ADL', '일상생활',
    '요실금', '인플루엔자', '폐렴구균',
    '위암', '대장암', '유방암', '자궁경부암',
]


def _is_data_row(ws, r):
    """A row is a data row if column 1 is an integer."""
    v = ws.cell(row=r, column=1).value
    return isinstance(v, (int, float))


def _find_header_row(ws):
    """Find the row that begins with '#' as the first cell."""
    for r in range(1, min(8, ws.max_row+1)):
        if ws.cell(row=r, column=1).value == '#':
            return r
    return None


def verify(path, golden=None, strict=False):
    errors, warnings = [], []
    wb = load_workbook(path)

    # 1. sheet count and order
    if wb.sheetnames != EXPECTED_SHEETS:
        errors.append(f'Sheet order/count mismatch. Got: {wb.sheetnames}')
        return errors, warnings, {}

    stats = {sn: {} for sn in EXPECTED_SHEETS}

    # 2-4: A and B sheets
    for sn in ('A_hc_finding', 'B_hc_clinical'):
        ws = wb[sn]
        hr = _find_header_row(ws)
        if hr is None:
            errors.append(f'[{sn}] header row (starts with "#") not found')
            continue
        actual_headers = [ws.cell(row=hr, column=c).value for c in range(1, 19)]
        if actual_headers != EXPECTED_HEADERS_18:
            errors.append(f'[{sn}] header mismatch at row {hr}.\n  Expected: {EXPECTED_HEADERS_18}\n  Got:      {actual_headers}')
        # data rows
        n_data, codes_used = 0, []
        for r in range(hr+1, ws.max_row+1):
            if not _is_data_row(ws, r): continue
            n_data += 1
            sc = ws.cell(row=r, column=7).value
            ps = ws.cell(row=r, column=9).value
            qc = ws.cell(row=r, column=10).value
            if not sc:
                errors.append(f'[{sn}] R{r}: snomed_code is empty')
            else:
                codes_used.append(str(sc))
            if ps not in ALLOWED_PS:
                errors.append(f'[{sn}] R{r}: presence_status invalid: {repr(ps)}')
            if ps in QUAL_MAP and qc != QUAL_MAP[ps]:
                warnings.append(f'[{sn}] R{r}: qualifier_code mismatch (ps={ps}, qc={qc}, expected {QUAL_MAP[ps]})')
        stats[sn]['data_rows'] = n_data
        stats[sn]['unique_codes'] = len(set(codes_used))

    # 5. always-generate in B sheet
    ws_b = wb['B_hc_clinical']
    hr_b = _find_header_row(ws_b)
    if hr_b is not None:
        items_b = []
        for r in range(hr_b+1, ws_b.max_row+1):
            if not _is_data_row(ws_b, r): continue
            items_b.append(ws_b.cell(row=r, column=6).value or '')
        text_blob = ' || '.join(items_b)
        for kw in ALWAYS_GENERATE_KEYWORDS:
            if kw not in text_blob:
                warnings.append(f'[B_hc_clinical] always-generate keyword missing: {kw}')

    # Basic sheet shape stats
    for sn in EXPECTED_SHEETS:
        ws = wb[sn]
        stats[sn]['max_row'] = ws.max_row
        stats[sn]['max_col'] = ws.max_column

    # 6. golden compare
    diff = {}
    if golden:
        gwb = load_workbook(golden)
        for sn in EXPECTED_SHEETS:
            if sn not in gwb.sheetnames:
                warnings.append(f'[golden] missing sheet: {sn}')
                continue
            mine = wb[sn].max_row
            ref = gwb[sn].max_row
            d = mine - ref
            diff[sn] = {'rows_mine': mine, 'rows_golden': ref, 'delta': d}
            if abs(d) > 5:
                msg = f'[{sn}] row count diverges: mine={mine}, golden={ref} (Δ={d:+d})'
                (errors if strict else warnings).append(msg)

    return errors, warnings, {'stats': stats, 'diff_vs_golden': diff}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('path', help='workbook to verify')
    ap.add_argument('--golden', help='reference workbook for row-count comparison')
    ap.add_argument('--strict', action='store_true', help='treat row-count divergence as error')
    args = ap.parse_args()

    errors, warnings, info = verify(args.path, golden=args.golden, strict=args.strict)
    out = {
        'path':     args.path,
        'errors':   errors,
        'warnings': warnings,
        'info':     info,
        'pass':     not errors,
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))
    if errors: return 2
    if warnings: return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
