#!/usr/bin/env python3
"""convert_to_v5 — Single-subject JSON → v5 5-sheet workbook.

Usage:
    python convert_to_v5.py --json INPUT.json --out OUTPUT.xlsx [--name 신정순]
    python convert_to_v5.py --xlsx FLAT132.xlsx --row 26 --out OUTPUT.xlsx
    python convert_to_v5.py --json INPUT.json --name 신정순 --compare-with PREV.json --out OUT.xlsx

The JSON may be either:
  (a) a single routineCheckupDataResponse dict, or
  (b) an envelope {"routineCheckupDataResponse": {...}}, or
  (c) a list of envelopes (e.g. the 83-subject file). With --name, picks that subject.

Exit code 0 on success, non-zero on error. Prints a small JSON report on stdout.
"""
from __future__ import annotations
import argparse, json, os, sys, traceback

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, THIS_DIR)

from v5_core import normalize_subject, derive_a_findings, derive_b_findings, derive_c_findings
from v5_writer import write_workbook


def _extract_rc(payload, name=None):
    """Find a routineCheckupDataResponse dict in the input."""
    if isinstance(payload, dict):
        if 'routineCheckupDataResponse' in payload:
            return payload['routineCheckupDataResponse']
        if any(k in payload for k in ('patientName', 'measurementTest', 'examination')):
            return payload
        return None
    if isinstance(payload, list):
        for entry in payload:
            rc = entry.get('routineCheckupDataResponse') if isinstance(entry, dict) else None
            if not rc: continue
            if name is None or rc.get('patientName') == name:
                return rc
    return None


def _load_from_xlsx(path, row):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
    flat = {headers[i]: ws.cell(row=row, column=i+1).value for i in range(len(headers))}
    # Re-shape flat 132-col row into routineCheckupDataResponse-like dict
    rc = {
        'patientName':       flat.get('patientName'),
        'birthDate':         flat.get('birthDate'),
        'examinationDate':   flat.get('examinationDate'),
        'organizationName':  flat.get('organizationName'),
        'screeningDate':     flat.get('screeningDate'),
        'suspectedDisease':  flat.get('suspectedDisease'),
        'existingDisease':   flat.get('existingDisease'),
        'additionalNotes':   flat.get('additionalNotes'),
        'comprehensiveOpinion': flat.get('comprehensiveOpinion'),
        'boneDensityTest':   flat.get('boneDensityTest'),
        'depression':        flat.get('depression'),
        'cognitiveDysfunction': flat.get('cognitiveDysfunction'),
        'geriatricPhysicalFunctionTest': flat.get('geriatricPhysicalFunctionTest'),
    }
    rc['measurementTest'] = {k.replace('measurementTest_',''): v
                             for k,v in flat.items()
                             if k and k.startswith('measurementTest_')}
    rc['urinaryTest']     = {'urinaryProtein': flat.get('urinaryTest_urinaryProtein')}
    rc['imagingTest']     = {'chestXray':      flat.get('imagingTest_chestXray')}
    rc['examination']     = {
        'pastMedicalHistory':  flat.get('examination_pastMedicalHistory'),
        'medicationTreatment': flat.get('examination_medicationTreatment'),
        'lifestyleFactors':    _reconstruct_lifestyle(flat),
    }
    rc['hepatitisB'] = {
        'hepatitisBEligibility': flat.get('hepatitisB_hepatitisBEligibility'),
        'surfaceAntigen': flat.get('hepatitisB_surfaceAntigen'),
        'surfaceAntibody': flat.get('hepatitisB_surfaceAntibody'),
        'testResult': flat.get('hepatitisB_testResult'),
    }
    rc['geriatricFunctionTest'] = {
        'geriatricEligibility': flat.get('geriatricFunctionTest_geriatricEligibility'),
        'fallHistory': flat.get('geriatricFunctionTest_fallHistory'),
        'adlScore': flat.get('geriatricFunctionTest_adlScore'),
        'urinaryIncontinence': flat.get('geriatricFunctionTest_urinaryIncontinence'),
        'vaccination': flat.get('geriatricFunctionTest_vaccination'),
    }
    # cancer screenings
    for kind in ('breastCancerScreenings','cervicalCancerScreenings','colorectalCancerScreenings','stomachCancerScreenings'):
        max_n = 4 if kind == 'colorectalCancerScreenings' else 2
        items = []
        for i in range(max_n):
            entry = {
                'biopsyResult':         flat.get(f'{kind}_{i}_biopsyResult'),
                'diagnosingDoctor':     flat.get(f'{kind}_{i}_diagnosingDoctor'),
                'recommendation':       flat.get(f'{kind}_{i}_recommendation'),
                'screeningAssessment':  flat.get(f'{kind}_{i}_screeningAssessment'),
                'screeningDate':        flat.get(f'{kind}_{i}_screeningDate'),
                'screeningInstitution': flat.get(f'{kind}_{i}_screeningInstitution'),
                'testItems':            flat.get(f'{kind}_{i}_testItems'),
                'testOpinion':          flat.get(f'{kind}_{i}_testOpinion'),
            }
            if any(v not in (None, '') for v in entry.values()):
                items.append(entry)
        rc[kind] = items
    return rc


def _reconstruct_lifestyle(flat):
    """Reverse the v3 split into the original ■/□ string."""
    parts = []
    for kw, key in (('금연','examination_lifestyleFactors_금연'),
                    ('절주','examination_lifestyleFactors_절주'),
                    ('신체활동','examination_lifestyleFactors_신체활동'),
                    ('근력운동','examination_lifestyleFactors_근력운동')):
        v = flat.get(key)
        if v == 'checked':         parts.append(f'■ {kw} 필요')
        elif v == 'not checked':   parts.append(f'□ {kw} 필요')
    return ' '.join(parts) if parts else None


def main():
    ap = argparse.ArgumentParser()
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument('--json', help='Input JSON file')
    src.add_argument('--xlsx', help='Input flattened 132-col xlsx file')
    ap.add_argument('--row', type=int, help='Row number (2-based) when using --xlsx')
    ap.add_argument('--name', help='Select subject by patientName when JSON contains a list')
    ap.add_argument('--compare-with', help='Optional second JSON for comparison mode (older exam)')
    ap.add_argument('--out', required=True, help='Output xlsx path')
    args = ap.parse_args()

    try:
        if args.json:
            with open(args.json, encoding='utf-8') as f:
                payload = json.load(f)
            rc = _extract_rc(payload, name=args.name)
            if rc is None:
                raise SystemExit(f'No routineCheckupDataResponse found (name filter: {args.name})')
        else:
            if not args.row:
                raise SystemExit('--row is required when using --xlsx')
            rc = _load_from_xlsx(args.xlsx, args.row)

        prev = None
        if args.compare_with:
            with open(args.compare_with, encoding='utf-8') as f:
                prev_payload = json.load(f)
            prev_rc = _extract_rc(prev_payload, name=args.name or rc.get('patientName'))
            if prev_rc:
                prev = normalize_subject(prev_rc)

        curr = normalize_subject(rc)
        a_rows = derive_a_findings(curr)
        b_rows = derive_b_findings(curr)
        c_rows = derive_c_findings(prev, curr) if prev else []

        os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
        write_workbook(args.out, curr, prev, a_rows, b_rows, c_rows)

        report = {
            'status':      'ok',
            'output':      args.out,
            'patient':     curr.get('patientName'),
            'patient_id':  curr.get('patient_id'),
            'exam_date':   curr.get('examinationDate'),
            'a_rows':      len(a_rows),
            'b_rows':      len(b_rows),
            'c_rows':      len(c_rows),
            'comparison':  bool(prev),
        }
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0
    except Exception as e:
        traceback.print_exc()
        print(json.dumps({'status': 'error', 'error': str(e)}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    sys.exit(main())
